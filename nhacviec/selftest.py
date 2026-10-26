"""Packaged-runtime checks use temporary data and do not import test runners."""
from datetime import date, datetime, time
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from . import __version__


def run_self_test(*, gui=False):
    from .models import Schedule
    from .storage import Storage
    from .excel import export_workbook
    from openpyxl import load_workbook

    checks = []
    with TemporaryDirectory(prefix="nhacviec-selftest-") as directory:
        base = Path(directory)
        with Storage(base / "data" / "tasks.db") as store:
            assert store.list_tasks() == []
            note = "Đang đối chiếu hồ sơ.\nChú ý mốc bàn giao và xác nhận số liệu."
            task = store.create_task("Công việc kiểm tra", datetime(2026, 10, 25, 16), category="Kiểm tra", notes=note)
            task2 = store.create_task("Việc chờ xử lý", datetime(2026, 10, 26, 9))
            assert task.id != task2.id and task.code != task2.code
            store.set_task_status(task.id, "in_progress")
            assert len(store.list_tasks(status="active")) == 2
            assert len(store.list_tasks(status="todo")) == 1
            checks.append("SQLite CRUD, unique codes and both task filters")
            assert store.get_task(task.id).notes == note
            assert [t.id for t in store.list_tasks(date_from=date(2026, 10, 25), date_to=date(2026, 10, 25))] == [task.id]
            checks.append("Unicode task notes and calendar deadline date lookup")

            store.save_schedule(Schedule(title="Lịch tháng ngày 31", start_date=date(2026, 1, 31), start_time=time(9), end_time=time(10), recurrence="monthly"))
            occurrences = store.occurrences(date(2026, 1, 1), date(2026, 3, 31))
            assert [o.start.date() for o in occurrences] == [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)]
            yearly = store.save_schedule(Schedule(title="Lịch năm ngày 29/2", start_date=date(2024, 2, 29), start_time=time(11), end_time=time(12), recurrence="yearly"))
            annual = [o for o in store.occurrences(date(2025, 1, 1), date(2025, 12, 31)) if o.schedule_id == yearly.id]
            assert [o.start.date() for o in annual] == [date(2025, 2, 28)]
            checks.append("Monthly/yearly recurrence with missing dates and preserved anchor")

            target = export_workbook(store, base / "report.xlsx", date(2026, 1, 1), date(2026, 3, 31))
            workbook = load_workbook(target)
            try:
                assert "Báo cáo tổng hợp" in workbook.sheetnames
                assert "DuLieuCongViec" in workbook["Báo cáo tổng hợp"].tables
                assert "DuLieuLichGoc" in workbook["Lịch gốc"].tables
                task_sheet = workbook["Báo cáo tổng hợp"]
                assert any(cell.value == note for row in task_sheet for cell in row)
            finally:
                workbook.close()
            checks.append("Excel workbook and named tables")

            store.backup(base / "backup.db")
            store.create_task("Chỉ có sau sao lưu", datetime(2026, 11, 1, 9))
            store.restore(base / "backup.db")
            assert len(store.list_tasks()) == 2
            assert store.get_task(task.id).notes == note
            checks.append("Consistent backup/restore")

            if gui:
                from .display import initialize_display, get_display_state
                initialize_display()
                import tkinter as tk
                from .app import NhacViecApp
                root = tk.Tk()
                root.withdraw()
                errors = []
                try:
                    app = NhacViecApp(root, store, base, integrations=False, enable_timers=False)
                    display = get_display_state(root)
                    if sys.platform == "win32" and getattr(sys, "frozen", False):
                        assert display["awareness"] == "per-monitor-v2", display
                        assert display["window_awareness"] == "per-monitor-v2", display
                    root.report_callback_exception = lambda *args: errors.append(str(args[1]))
                    root.deiconify()
                    root.update()
                    assert len(app.pages["active"].visible_tasks) == 2
                    app.show_page("calendar")
                    calendar = app.pages["calendar"]
                    calendar.anchor = calendar.selected_date = date(2026, 10, 25)
                    for mode in ("week", "month", "year"):
                        calendar.mode.set(mode)
                        calendar.refresh()
                        root.update()
                        assert [t.id for t in calendar.task_cards.visible_tasks] == [task.id]
                        assert calendar.task_cards.visible_tasks[0].notes == note
                    notes = app.edit_task_notes(task.id)
                    root.update()
                    notes.notes_text.insert("end", "\nĐã cập nhật trong cửa sổ ghi chú.")
                    notes.save()
                    root.update()
                    assert store.get_task(task.id).notes.endswith("Đã cập nhật trong cửa sổ ghi chú.")
                    assert calendar.task_cards.visible_tasks[0].notes == store.get_task(task.id).notes
                    app.show_page("active")
                    root.update()
                    assert not errors, errors
                    app.exit()
                    checks.append("Tk desktop, deadline task cards in week/month/year and saved notes refresh")
                finally:
                    try:
                        root.destroy()
                    except tk.TclError:
                        pass
    report = {"passed": True, "gui": gui, "version": __version__, "checks": checks}
    if gui:
        report["display"] = display
    return report
