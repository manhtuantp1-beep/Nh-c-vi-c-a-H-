from datetime import date, datetime, time
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from nhacviec.excel import export_workbook
from nhacviec.storage import Storage


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = Storage(Path(self.temp.name) / "tasks.db")
        self.target = Path(self.temp.name) / "report.xlsx"

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def test_export_all_tasks_bounded_occurrences_and_named_tables_without_mutating(self):
        self.store.create_task("Ngoài kỳ", datetime(2031, 1, 1))
        s = self.store.create_schedule(title="Lịch", start_date=date(2030, 1, 1), start_time=time(9), end_time=time(10), recurrence="monthly")
        self.store.set_occurrence_status(s.id, date(2030, 1, 1), "completed")
        count_before = self.store.connection.execute("SELECT count(*) FROM occurrence_overrides").fetchone()[0]
        export_workbook(self.store, self.target, date(2030, 1, 1), date(2030, 2, 1))
        self.assertEqual(self.store.connection.execute("SELECT count(*) FROM occurrence_overrides").fetchone()[0], count_before)
        wb = load_workbook(self.target)
        try:
            sheet = wb["Báo cáo tổng hợp"]
            self.assertIn("DuLieuCongViec", sheet.tables)
            self.assertIn("DuLieuLichGoc", wb["Lịch gốc"].tables)
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(sheet.cell(2, 3).value, "Ngoài kỳ")
            self.assertEqual({sheet.cell(row, 16).value for row in (3, 4)}, {"Đã lưu", "Dự kiến"})
            self.assertEqual(sheet.cell(2, 6).value, datetime(2031, 1, 1))
            self.assertIsInstance(sheet.cell(2, 21).value, datetime)
        finally:
            wb.close()

    def test_formula_like_input_is_preserved_as_text(self):
        self.store.create_task("=1+1", datetime(2030, 1, 1), content="=HYPERLINK(\"x\")")
        export_workbook(self.store, self.target, date(2030, 1, 1), date(2030, 1, 1))
        wb = load_workbook(self.target)
        try:
            self.assertEqual(wb.active.cell(2, 3).data_type, "s")
            self.assertEqual(wb.active.cell(2, 3).value, "=1+1")
            self.assertEqual(wb.active.cell(2, 9).data_type, "s")
        finally:
            wb.close()

    def test_empty_workbook_has_headers_and_tables_without_fake_records(self):
        export_workbook(self.store, self.target, date(2030, 1, 1), date(2030, 1, 1))
        wb = load_workbook(self.target)
        try:
            self.assertEqual(wb.active.max_row, 1)
            self.assertEqual(wb["Lịch gốc"].max_row, 1)
            self.assertEqual(wb.active.tables["DuLieuCongViec"].ref, "A1:U1")
        finally:
            wb.close()

    def test_invalid_range_and_database_destination_rejected(self):
        with self.assertRaises(ValueError):
            export_workbook(self.store, self.target, date(2030, 2, 1), date(2030, 1, 1))
        with self.assertRaises(ValueError):
            export_workbook(self.store, self.store.path, date(2030, 1, 1), date(2030, 1, 1))


if __name__ == "__main__":
    unittest.main()
