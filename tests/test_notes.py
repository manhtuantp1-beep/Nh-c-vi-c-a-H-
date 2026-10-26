"""Progress notes, existing-database upgrade and non-destructive restore tests."""
from contextlib import closing
from dataclasses import asdict
from datetime import date, datetime, time
from pathlib import Path
import sqlite3
import tempfile
import tkinter as tk
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from nhacviec.dialogs import NotesDialog, TaskDialog
from nhacviec.excel import export_workbook
from nhacviec.storage import SCHEMA_VERSION, Storage, UnsupportedDatabase
from nhacviec.theme import configure_style


def make_v1(path):
    """Construct the original app schema with real task/calendar/history rows."""
    with Storage(path) as store:
        task = store.create_task("Hồ sơ đang làm", datetime(2030, 1, 8, 17),
                                 category="Pháp lý", content="Nội dung\nĐầy đủ", requirements="Bản gốc")
        task = store.set_task_status(task.id, "in_progress", "Đang rà soát")
        schedule = store.create_schedule(title="Họp", start_date=date(2030, 1, 8),
                                         start_time=time(9), end_time=time(10), recurrence="once")
        store.set_occurrence_status(schedule.id, date(2030, 1, 8), "completed")
        store.save_settings(days_before=14)
        store.mark_notification("giữ-lịch-sử", datetime(2030, 1, 8, 8))
    with closing(sqlite3.connect(path)) as connection, connection:
        connection.execute("ALTER TABLE tasks DROP COLUMN notes")
        connection.execute("PRAGMA user_version=1")
    return task, schedule


class NotesStorageTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "tasks.db"

    def tearDown(self):
        self.directory.cleanup()

    def test_notes_survive_lifecycle_reopen_backup_and_restore(self):
        text = "Đã gửi hồ sơ cho anh Hùng.\nChú ý: kiểm tra số liệu trước 16 giờ.\n"
        with Storage(self.path) as store:
            task = store.create_task("Báo cáo", datetime(2030, 1, 8), content="Nội dung riêng",
                                     requirements="Yêu cầu riêng", notes=text)
            for state in ("in_progress", "completed", "todo", "canceled", "todo"):
                store.set_task_status(task.id, state)
                self.assertEqual(store.get_task(task.id).notes, text)
            backup = store.backup(self.path.with_name("notes-backup.db"))
            store.update_task(task.id, notes="")
            self.assertEqual(store.get_task(task.id).notes, "")
            before_restore = store.restore(backup)
            self.assertEqual(store.get_task(task.id).notes, text)
            with Storage(before_restore) as previous:
                self.assertEqual(previous.get_task(task.id).notes, "")
        with Storage(self.path) as store:
            saved = store.get_task(task.id)
            self.assertEqual((saved.notes, saved.content, saved.requirements),
                             (text, "Nội dung riêng", "Yêu cầu riêng"))

    def test_v1_upgrade_keeps_task_calendar_settings_and_notification_data(self):
        task, schedule = make_v1(self.path)
        with Storage(self.path) as store:
            self.assertEqual(store.connection.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
            self.assertEqual(asdict(store.get_task(task.id)), asdict(task))
            self.assertEqual(store.get_task(task.id).notes, "")
            self.assertEqual(store.get_schedule(schedule.id), schedule)
            self.assertEqual(store.occurrences(date(2030, 1, 8), date(2030, 1, 8))[0].status, "completed")
            self.assertEqual(store.get_settings()["days_before"], 14)
            self.assertTrue(store.notification_delivered("giữ-lịch-sử"))
            store.update_task(task.id, notes="Ghi chú mới sau nâng cấp")
        backups = list(self.path.parent.glob("tasks.pre-upgrade-*.db"))
        self.assertEqual(len(backups), 1)
        with closing(sqlite3.connect(backups[0])) as old:
            self.assertEqual(old.execute("PRAGMA user_version").fetchone()[0], 1)
            self.assertNotIn("notes", {row[1] for row in old.execute("PRAGMA table_info(tasks)")})
            self.assertEqual(old.execute("SELECT title FROM tasks").fetchone()[0], task.title)
        with Storage(self.path) as store:
            self.assertEqual(store.get_task(task.id).notes, "Ghi chú mới sau nâng cấp")
        self.assertEqual(list(self.path.parent.glob("tasks.pre-upgrade-*.db")), backups)

    def test_restore_v1_backup_migrates_copy_and_preserves_source_and_live_backup(self):
        source = self.path.with_name("old-backup.db")
        task, _schedule = make_v1(source)
        source_bytes = source.read_bytes()
        with Storage(self.path) as store:
            existing = store.create_task("Trước khôi phục", datetime(2030, 2, 1), notes="Ghi chú đang có")
            saved = store.restore(source)
            restored = store.get_task(task.id)
            self.assertEqual((restored.title, restored.notes), (task.title, ""))
            store.update_task(restored.id, notes="Đã nhận đủ hồ sơ")
            self.assertEqual(store.connection.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
            with Storage(saved) as previous:
                self.assertEqual(previous.get_task(existing.id).notes, "Ghi chú đang có")
        self.assertEqual(source.read_bytes(), source_bytes)
        with closing(sqlite3.connect(source)) as old:
            self.assertEqual(old.execute("PRAGMA user_version").fetchone()[0], 1)

    def test_failed_staged_migration_leaves_live_data_and_source_unchanged(self):
        source = self.path.with_name("old-backup.db")
        make_v1(source)
        source_bytes = source.read_bytes()
        with Storage(self.path) as store:
            existing = store.create_task("Giữ nguyên", datetime(2030, 2, 1), notes="Không mất ghi chú")
            with patch.object(Storage, "_migrate_database", side_effect=sqlite3.OperationalError("disk full")):
                with self.assertRaises(sqlite3.OperationalError):
                    store.restore(source)
            self.assertEqual(store.get_task(existing.id).notes, "Không mất ghi chú")
            self.assertEqual(store.connection.execute("PRAGMA user_version").fetchone()[0], SCHEMA_VERSION)
        self.assertEqual(source.read_bytes(), source_bytes)

    def test_add_column_and_version_change_rollback_together(self):
        make_v1(self.path)
        with closing(sqlite3.connect(self.path)) as connection:
            def authorizer(action, first, second, _database, _trigger):
                if action == sqlite3.SQLITE_PRAGMA and first == "user_version" and second == str(SCHEMA_VERSION):
                    return sqlite3.SQLITE_DENY
                return sqlite3.SQLITE_OK
            connection.set_authorizer(authorizer)
            with self.assertRaises(sqlite3.DatabaseError):
                Storage._migrate_database(connection)
            connection.set_authorizer(None)
            self.assertEqual(connection.execute("PRAGMA user_version").fetchone()[0], 1)
            self.assertNotIn("notes", {row[1] for row in connection.execute("PRAGMA table_info(tasks)")})

    def test_foreign_broken_or_future_databases_are_not_migrated(self):
        for invalid in ("foreign", "broken", "future", "mismatched"):
            with self.subTest(invalid=invalid):
                path = self.path.with_name(f"{invalid}.db")
                make_v1(path)
                with closing(sqlite3.connect(path)) as connection, connection:
                    if invalid == "foreign":
                        connection.execute("UPDATE app_metadata SET value='other-app' WHERE key='app_id'")
                    elif invalid == "broken":
                        connection.execute("UPDATE tasks SET deadline=''")
                    elif invalid == "future":
                        connection.execute("PRAGMA user_version=99")
                    else:
                        connection.execute("PRAGMA user_version=2")
                before = path.read_bytes()
                with self.assertRaises(UnsupportedDatabase):
                    Storage(path)
                self.assertEqual(path.read_bytes(), before)
                self.assertEqual(list(path.parent.glob(f"{invalid}.pre-upgrade-*.db")), [])

    def test_invalid_note_types_leave_existing_note_unchanged(self):
        with Storage(self.path) as store:
            task = store.create_task("Công việc", datetime(2030, 1, 8), notes="Giữ lại")
            with self.assertRaises(ValueError):
                store.update_task(task.id, notes=None)
            with self.assertRaises(ValueError):
                store.create_task("Sai", datetime(2030, 1, 8), notes=123)
            self.assertEqual(store.get_task(task.id).notes, "Giữ lại")
            self.assertEqual(len(store.list_tasks()), 1)

    def test_excel_exports_multiline_notes_as_text_and_keeps_schedule_column_blank(self):
        text = '=HYPERLINK("x")\nGhi chú tiếng Việt đầy đủ'
        target = self.path.with_suffix(".xlsx")
        with Storage(self.path) as store:
            task = store.create_task("Công việc", datetime(2030, 1, 8), notes=text)
            store.create_schedule(title="Họp", start_date=date(2030, 1, 8), start_time=time(9),
                                   end_time=time(10), recurrence="once")
            export_workbook(store, target, date(2030, 1, 8), date(2030, 1, 8))
            self.assertEqual(store.get_task(task.id).notes, text)
        workbook = load_workbook(target)
        try:
            sheet = workbook.active
            column = next(cell.column for cell in sheet[1] if cell.value == "Ghi chú công việc")
            self.assertEqual(sheet.cell(2, column).value, text)
            self.assertEqual(sheet.cell(2, column).data_type, "s")
            self.assertIsNone(sheet.cell(3, column).value)
            self.assertEqual(sheet.tables["DuLieuCongViec"].ref, "A1:V3")
        finally:
            workbook.close()


class NotesDialogTests(unittest.TestCase):
    def setUp(self):
        try:
            self.root = tk.Tk()
        except tk.TclError as exc:
            self.skipTest(f"Tk display unavailable: {exc}")
        configure_style(self.root)
        self.directory = tempfile.TemporaryDirectory()
        self.store = Storage(Path(self.directory.name) / "tasks.db")
        self.app = SimpleNamespace(root=self.root, store=self.store, refresh_current=Mock())
        self.task = self.store.create_task("Rà soát báo cáo", datetime(2030, 1, 8),
                                           content="Nội dung gốc", requirements="Yêu cầu gốc", notes="Đã nhận dữ liệu")

    def tearDown(self):
        if hasattr(self, "root"):
            self.root.destroy()
            self.store.close()
            self.directory.cleanup()

    def test_quick_editor_saves_notes_without_overwriting_other_task_fields(self):
        dialog = NotesDialog(self.app, self.task)
        self.assertEqual(dialog.snapshot(), "Đã nhận dữ liệu")
        dialog.notes_text.insert("end", "\nChờ đối chiếu lần cuối")
        self.store.update_task(self.task.id, title="Tên mới", status="in_progress")
        dialog.save()
        saved = self.store.get_task(self.task.id)
        self.assertEqual((saved.title, saved.status, saved.content, saved.requirements),
                         ("Tên mới", "in_progress", "Nội dung gốc", "Yêu cầu gốc"))
        self.assertEqual(saved.notes, "Đã nhận dữ liệu\nChờ đối chiếu lần cuối")
        self.app.refresh_current.assert_called_once()

    def test_quick_editor_cancel_honors_unsaved_notes(self):
        dialog = NotesDialog(self.app, self.task)
        dialog.notes_text.insert("end", "\nChưa lưu")
        with patch("nhacviec.dialogs.messagebox.askyesno", return_value=False) as prompt:
            dialog.cancel()
            prompt.assert_called_once()
        self.assertTrue(dialog.winfo_exists())
        self.assertEqual(self.store.get_task(self.task.id).notes, "Đã nhận dữ liệu")
        with patch("nhacviec.dialogs.messagebox.askyesno", return_value=True):
            dialog.cancel()
        self.assertFalse(dialog.winfo_exists())
        self.app.refresh_current.assert_not_called()

    def test_main_task_editor_includes_notes_and_preserves_multiline_text(self):
        dialog = TaskDialog(self.app, self.task)
        self.assertEqual(dialog.notes_text.get("1.0", "end-1c"), self.task.notes)
        dialog.notes_text.delete("1.0", "end")
        dialog.notes_text.insert("1.0", "Đã làm xong phần A.\nCần hỏi lại phần B.")
        dialog.save()
        self.assertEqual(self.store.get_task(self.task.id).notes, "Đã làm xong phần A.\nCần hỏi lại phần B.")
        self.app.refresh_current.assert_called_once()

    def test_main_task_editor_dirty_guard_includes_notes(self):
        dialog = TaskDialog(self.app, self.task)
        dialog.notes_text.insert("end", "\nCần lưu")
        with patch("nhacviec.dialogs.messagebox.askyesno", return_value=False) as prompt:
            dialog.cancel()
            prompt.assert_called_once()
        self.assertTrue(dialog.winfo_exists())
        self.assertEqual(self.store.get_task(self.task.id).notes, self.task.notes)
        dialog.destroy()


if __name__ == "__main__":
    unittest.main()
